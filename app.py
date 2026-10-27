import io, os, re, uuid, zipfile, unicodedata, datetime as dt
from flask import Flask, request, jsonify, render_template, send_file, abort
from openpyxl import load_workbook
from docfill import fill
from Codebarre import inserer_codes_barres, repere
from mise_en_page import aligner, accorder
import traduction_ar as TR
import genre

app = Flask(__name__)
BASE = os.path.dirname(os.path.abspath(__file__))
UP = os.path.join(BASE, 'uploads'); MOD = os.path.join(BASE, 'modeles')
PROGRAMMES = {'doctorat': 'Doctorat', 'habilitation': 'Habilitation universitaire (brouillon)'}
DOCS_VALIDES = ('pv', 'diplome', 'attestation', 'invitation', 'resume')
JOURS = ['Lundi','Mardi','Mercredi','Jeudi','Vendredi','Samedi','Dimanche']
MOIS = ['Janvier','Février','Mars','Avril','Mai','Juin','Juillet','Août','Septembre','Octobre','Novembre','Décembre']
os.makedirs(UP, exist_ok=True)
BLANC = '………'
MOIS_AR = ['يناير','فبراير','مارس','أبريل','ماي','يونيو','يوليوز','غشت','شتنبر','اكتوبر','نونبر','دجنبر']
GRADE_AR = {'PES': 'أستاذ التعليم العالي', 'PH': 'أستاذ مؤهل', 'MCH': 'أستاذ محاضر مؤهل', 'MC': 'أستاذ محاضر', 'P': 'أستاذ'}
GRADE_AR_F = {'PES': 'أستاذة التعليم العالي', 'PH': 'أستاذة مؤهلة', 'MCH': 'أستاذة محاضرة مؤهلة', 'MC': 'أستاذة محاضرة', 'P': 'أستاذة'}
def grade_ar(code, sexe=''): return (GRADE_AR_F if sexe == 'F' else GRADE_AR).get(code, '')
ETAB_DEFAUT, ETAB_AR_DEFAUT = 'Faculté des Sciences, Agadir', 'كلية العلوم بأگادير'
CED_DEFAUT, CED_AR_DEFAUT = 'Sciences et Techniques et Sciences Médicales', 'العلوم والتقنيات والعلوم الطبية'
DISCIPLINE_AR = {'informatique': 'الإعلاميات', 'biologie': 'البيولوجيا', 'chimie': 'الكيمياء', 'physique': 'الفيزياء',
                 'mathematiques': 'الرياضيات', 'geologie': 'الجيولوجيا', 'mathematiques appliquees': 'الرياضيات التطبيقية'}
NB_JURY = 10
INST_AR = {'faculte des sciences': 'كلية العلوم', 'faculte des sciences et techniques': 'كلية العلوم والتقنيات',
           'faculte polydisciplinaire': 'الكلية المتعددة التخصصات', 'ecole superieure de technologie': 'المدرسة العليا للتكنولوجيا',
           "ecole nationale superieure d'informatique et d'analyse des systemes": 'المدرسة الوطنية العليا للمعلوميات وتحليل النظم'}
VILLE_AR = {'agadir': 'أگادير', 'taroudant': 'تارودانت', 'rabat': 'الرباط', 'marrakech': 'مراكش', 'casablanca': 'الدار البيضاء',
            'meknes': 'مكناس', 'fes': 'فاس', 'kenitra': 'القنيطرة', 'tetouan': 'تطوان', 'oujda': 'وجدة', 'settat': 'سطات',
            'ouarzazate': 'ورزازات', 'essaouira': 'الصويرة', 'beni mellal': 'بني ملال', 'el jadida': 'الجديدة', 'tanger': 'طنجة',
            'safi': 'آسفي', 'laayoune': 'العيون', 'guelmim': 'كلميم', 'ait melloul': 'آيت ملول', 'tiznit': 'تيزنيت',
            'errachidia': 'الرشيدية', 'inezgane': 'إنزكان'}
def etab_ar(parts):
    """['Faculté des Sciences', 'Agadir'] -> 'كلية العلوم ، أگادير' ; texte d'origine si l'établissement n'est pas dans les dictionnaires."""
    if len(parts) >= 2:
        inst = INST_AR.get(sans_accents(', '.join(parts[:-1])).lower().replace('’', "'"))
        ville = VILLE_AR.get(sans_accents(parts[-1]).lower())
        if inst and ville: return f'{inst} ، {ville}'
    return ' ، '.join(parts)
MENTION_AR = {'très honorable': 'مشرف جدا', 'honorable': 'مشرف'}
FORMATION_AR = {'sciences et techniques & sciences médicales': 'العلوم والتقنيات والعلوم الطبية'}


def txt(v): return '' if v is None else str(v).strip()
def date_fr(v): return v.strftime('%d/%m/%Y') if isinstance(v, (dt.date, dt.datetime)) else txt(v)
def date_ar(v): return f"{v.day:02d} {MOIS_AR[v.month-1]} {v.year}" if isinstance(v, (dt.date, dt.datetime)) else txt(v)
def date_longue(v): return f"{v.day:02d} {MOIS[v.month-1].lower()} {v.year}" if isinstance(v, (dt.date, dt.datetime)) else ''
def code39(s):
    """Texte pour la police « Barcode 3 of 9 » : *CODE* avec les caractères au format Word (U+F000 + ASCII)."""
    s = re.sub(r'[^0-9A-Z\-. $/+%]', '', txt(s).upper())
    return ''.join(chr(0xF000 + ord(c)) for c in '*' + s + '*') if s else ''


def propre(s):
    """Nom sûr pour un chemin dans le zip (pas de / \\ : * ? " < > |)."""
    s = re.sub(r'[\\/:*?"<>|]+', '_', txt(s))
    return re.sub(r'\s+', '_', s) or 'inconnu'


def norm_col(c):
    """'Prénom ' -> 'PRENOM', 'Date naissance' -> 'DATE_NAISSANCE' (sans accents, espaces -> _)."""
    s = unicodedata.normalize('NFD', txt(c))
    s = ''.join(ch for ch in s if unicodedata.category(ch) != 'Mn')
    return re.sub(r'[^A-Z0-9]+', '_', s.upper()).strip('_')


def colonne_specialite(c):
    """« SPECIALITE », « SPETIALITE », « SPECIALITE_DU_CANDIDAT », « SPECIALITES »… -> 'SPECIALITE' (ou 'SPECIALITE_AR').
    La spécialité des rapporteurs / membres du jury n'est pas concernée."""
    if not re.match(r'^SPE[A-Z]{0,3}IA?LITES?(_|$)', c) or re.search(r'RAPPORTEUR|JURY|MEMBRE', c): return None
    return 'SPECIALITE_AR' if 'ARABE' in c or c.endswith('_AR') else 'SPECIALITE'


def colonne_connue(c):
    return colonne_specialite(c) or c


def lire(path):
    """Retourne une liste de dict ; les en-têtes sont normalisés (majuscules, sans accents)."""
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        rows = list(wb.active.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows: return []
    cols = [colonne_connue(norm_col(c)) for c in rows[0]]
    data = []
    for r in rows[1:]:
        if not any(x not in (None, '') for x in r): continue
        d = {}
        for c, x in zip(cols, r):
            if c and (x not in (None, '') or c not in d): d[c] = x      # une colonne vide n'écrase pas une colonne remplie du même nom
        data.append(d)
    if 'HORODATEUR' in cols:                    # export du formulaire (Doctorants_*.xlsx)
        data = [depuis_formulaire(d) for d in data]
    return [enrichir_ar(d) for d in data]


# ------------------------------------------------------------------ traduction de secours (si traduction_ar ne trouve rien)
TERMES_AR = {   # expressions complètes (clé : minuscules, sans accents)
    'mathematiques, informatique et applications': 'الرياضيات والإعلاميات والتطبيقات',
    'securite informatique': 'الأمن المعلوماتي', 'intelligence artificielle': 'الذكاء الاصطناعي',
    'sciences de la vie': 'علوم الحياة', 'sciences de la terre': 'علوم الأرض', 'sciences des donnees': 'علوم المعطيات',
    'sciences de la vie et ressources naturelles': 'علوم الحياة والموارد الطبيعية',
    'chimie environnementale et valorisation des dechets agroalimentaires': 'الكيمياء البيئية وتثمين النفايات الغذائية',
    'valorisation des dechets agroalimentaires': 'تثمين النفايات الغذائية',
    'materiaux, thermodynamique et energetique': 'المواد والديناميكا الحرارية والطاقيات',
    'chimie physique': 'الكيمياء الفيزيائية', 'chimie theorique': 'الكيمياء النظرية', 'chimie organique': 'الكيمياء العضوية',
    'chimie analytique': 'الكيمياء التحليلية', 'chimie inorganique': 'الكيمياء غير العضوية',
    'informatique et reseaux': 'الإعلاميات والشبكات', 'genie informatique': 'الهندسة المعلوماتية',
    'genie civil': 'الهندسة المدنية', 'genie des procedes': 'هندسة الطرائق', 'energies renouvelables': 'الطاقات المتجددة',
    'sciences et techniques et sciences medicales': 'العلوم والتقنيات والعلوم الطبية',
    'sciences et techniques & sciences medicales': 'العلوم والتقنيات والعلوم الطبية',
    'analyse fonctionnelle': 'التحليل الدالي', 'analyse numerique': 'التحليل العددي', 'analyse harmonique': 'التحليل التوافقي',
    'recherche operationnelle': 'بحوث العمليات', 'equations aux derivees partielles': 'المعادلات التفاضلية الجزئية',
    'physique des materiaux': 'فيزياء المواد', 'physique du solide': 'فيزياء الحالة الصلبة', 'physique theorique': 'الفيزياء النظرية',
    'traitement d\'images': 'معالجة الصور', 'apprentissage automatique': 'التعلم الآلي', 'systemes d\'information': 'نظم المعلومات',
    'tres honorable': 'مشرف جدا', 'honorable': 'مشرف'}
NOMS_AR = {     # noms (avec l'article)
    'mathematiques': 'الرياضيات', 'mathematique': 'الرياضيات', 'informatique': 'الإعلاميات', 'informatiques': 'الإعلاميات',
    'applications': 'التطبيقات', 'biologie': 'البيولوجيا', 'physique': 'الفيزياء', 'chimie': 'الكيمياء', 'geologie': 'الجيولوجيا',
    'sciences': 'العلوم', 'science': 'العلوم', 'vie': 'الحياة', 'terre': 'الأرض', 'ressources': 'الموارد', 'materiaux': 'المواد',
    'thermodynamique': 'الديناميكا الحرارية', 'energetique': 'الطاقيات', 'energie': 'الطاقة', 'reseaux': 'الشبكات',
    'securite': 'الأمن', 'environnement': 'البيئة', 'electronique': 'الإلكترونيات', 'mecanique': 'الميكانيك',
    'statistique': 'الإحصاء', 'statistiques': 'الإحصاء', 'biotechnologie': 'البيوتكنولوجيا', 'ecologie': 'الإيكولوجيا',
    'microbiologie': 'الميكروبيولوجيا', 'biochimie': 'الكيمياء الحيوية', 'genie': 'الهندسة', 'optique': 'البصريات',
    'donnees': 'المعطيات', 'telecommunications': 'الاتصالات', 'valorisation': 'التثمين', 'dechets': 'النفايات',
    'modelisation': 'النمذجة', 'analyse': 'التحليل', 'algebre': 'الجبر', 'geometrie': 'الهندسة', 'techniques': 'التقنيات',
    'technologie': 'التكنولوجيا', 'eau': 'الماء', 'agronomie': 'الهندسة الزراعية', 'nutrition': 'التغذية', 'sante': 'الصحة',
    'electrotechnique': 'الكهروتقنية', 'automatique': 'الآلية', 'robotique': 'الروبوتيك', 'genetique': 'علم الوراثة',
    'physiologie': 'الفيزيولوجيا', 'botanique': 'علم النبات', 'zoologie': 'علم الحيوان', 'oceanographie': 'علم المحيطات',
    'aquaculture': 'تربية الأحياء المائية', 'materiaux': 'المواد', 'catalyse': 'الحفز', 'corrosion': 'التآكل'}
ADJ_AR = {      # adjectifs (forme féminine, valable pour les noms ci-dessus au pluriel non humain)
    'appliquees': 'التطبيقية', 'appliquee': 'التطبيقية', 'applique': 'التطبيقي', 'naturelles': 'الطبيعية', 'naturelle': 'الطبيعية',
    'physique': 'الفيزيائية', 'theorique': 'النظرية', 'environnementale': 'البيئية', 'environnementales': 'البيئية',
    'informatique': 'المعلوماتية', 'artificielle': 'الاصطناعية', 'medicales': 'الطبية', 'medicale': 'الطبية',
    'agroalimentaires': 'الغذائية', 'agroalimentaire': 'الغذائية', 'organique': 'العضوية', 'analytique': 'التحليلية',
    'renouvelables': 'المتجددة', 'fondamentales': 'الأساسية', 'fondamentale': 'الأساسية', 'numerique': 'الرقمية',
    'marine': 'البحرية', 'marines': 'البحرية', 'vegetale': 'النباتية', 'animale': 'الحيوانية', 'moleculaire': 'الجزيئية',
    'cellulaire': 'الخلوية', 'nucleaire': 'النووية', 'quantique': 'الكمية', 'statistique': 'الإحصائية', 'electrique': 'الكهربائية'}
LIEUX_AR = dict(VILLE_AR, **{
    'ida outanane': 'إداوتنان', 'agadir ida outanane': 'أگادير إداوتنان', 'ouled teima': 'أولاد تايمة', 'tata': 'طاطا',
    'youssoufia': 'اليوسفية', 'sidi slimane': 'سيدي سليمان', 'tunis': 'تونس', 'boujaad': 'أبي الجعد', 'tinghir': 'تنغير',
    'drargua': 'الدراركة', 'tissint': 'تسينت', 'khouribga': 'خريبكة', 'nador': 'الناظور', 'taza': 'تازة', 'larache': 'العرائش',
    'khemisset': 'الخميسات', 'berkane': 'بركان', 'dakhla': 'الداخلة', 'tan tan': 'طانطان', 'chtouka ait baha': 'اشتوكة آيت باها',
    'biougra': 'بيوكرى', 'dcheira': 'الدشيرة', 'mohammedia': 'المحمدية', 'sale': 'سلا', 'temara': 'تمارة', 'ifrane': 'إفران',
    'lagfifat': 'لكفيفات', 'tagadirt': 'تاكاديرت', 'dr': 'دوار', 'douar': 'دوار', 'ikniouen': 'إكنيون',
    'azrou': 'أزرو', 'khenifra': 'خنيفرة', 'zagora': 'زاكورة', 'sidi ifni': 'سيدي إفني', 'tafraout': 'تافراوت'})


def cle_ar(s):
    return re.sub(r'\s+', ' ', sans_accents(s).lower().replace('’', "'")).strip(' .')


def _partie_ar(p):
    """Un morceau sans « et » : expression connue, « X de Y » (idafa) ou nom + adjectifs. '' si un mot est inconnu."""
    if p in TERMES_AR: return TERMES_AR[p]
    if p in NOMS_AR: return NOMS_AR[p]
    m = re.match(r"^(.+?) (?:de la|de l'|des|du|de|d')\s*(.+)$", p)
    if m:
        a, b = _partie_ar(m.group(1)), _partie_ar(m.group(2))
        if a and b:
            mots = a.split(' ', 1); tete = mots[0][2:] if mots[0].startswith('ال') else mots[0]
            return ' '.join([tete] + mots[1:] + [b]) if len(mots) == 1 else f'{tete} {b} {mots[1]}'
        return ''
    mots = p.split()
    if not mots or mots[0] not in NOMS_AR: return ''
    adj = [ADJ_AR.get(w) for w in mots[1:]]
    return ' '.join([NOMS_AR[mots[0]]] + adj) if all(adj) else ''


def traduire_terme(s):
    """Spécialité / discipline / formation / mention en français -> arabe ('' si la traduction n'est pas sûre)."""
    c = cle_ar(s)
    if not c: return ''
    if c in TERMES_AR: return TERMES_AR[c]
    parties = [x for x in re.split(r"\s*(?:,|;|/|&|\bet\b|\s-\s|-|–|—)\s*", c) if x]
    ar = [_partie_ar(x) for x in parties]
    return ' و'.join(ar) if ar and all(ar) else ''


def traduire_lieu(s):
    """« Agadir Ida Outanane », « Tissint Tata » -> arabe, mot à mot ; '' si un nom de lieu est inconnu."""
    mots = re.sub(r'[,\-]', ' ', cle_ar(s)).split(); out, i = [], 0
    while i < len(mots):
        for n in (3, 2, 1):
            g = ' '.join(mots[i:i + n])
            if len(mots[i:i + n]) == n and g in LIEUX_AR: out.append(LIEUX_AR[g]); i += n; break
        else: return ''
    return ' '.join(out)


def _tr(nom, *a):
    """Appel protégé de traduction_ar.<nom> (fonction absente ou erreur -> '')."""
    try: return txt(getattr(TR, nom)(*a))
    except Exception: return ''


def enrichir_ar(d):
    """Complète en arabe (traduction automatique) ce que l'Excel ne donne pas ; ne remplace jamais une valeur saisie."""
    tr = []
    def mettre(cle_ar, valeur, nom):
        if not txt(d.get(cle_ar)) and valeur: d[cle_ar] = valeur; tr.append(nom)
    if not txt(d.get('SPECIALITE')) and txt(d.get('DISCIPLINE')):       # pas de colonne « Spécialité » : on reprend la discipline
        d['SPECIALITE'] = d['DISCIPLINE']; d['SPECIALITE_DEDUITE'] = 'oui'
    if not txt(d.get('SPECIALITE_AR')) and txt(d.get('DISCIPLINE_AR')) and d.get('SPECIALITE_DEDUITE'):
        d['SPECIALITE_AR'] = d['DISCIPLINE_AR']
    mettre('LIEU_NAISSANCE_AR', _tr('lieu', d.get('LIEU_NAISSANCE')) or traduire_lieu(d.get('LIEU_NAISSANCE')), 'lieu de naissance')
    mettre('DISCIPLINE_AR', _tr('discipline', d.get('DISCIPLINE')) or traduire_terme(d.get('DISCIPLINE')), 'discipline')
    mettre('SPECIALITE_AR', _tr('specialite', d.get('SPECIALITE')) or traduire_terme(d.get('SPECIALITE')) or
           (txt(d.get('DISCIPLINE_AR')) if d.get('SPECIALITE_DEDUITE') else ''), 'spécialité')
    mettre('MENTION_AR', _tr('mention', d.get('MENTION')) or traduire_terme(d.get('MENTION')), 'mention')
    mettre('FORMATION_AR', _tr('formation', d.get('FORMATION')) or traduire_terme(d.get('FORMATION')), 'formation')
    mettre('CED_AR', traduire_terme(d.get('CED')), 'CED')
    if not txt(d.get('NOM_AR')) and not txt(d.get('PRENOM_AR')): mettre('NOM_AR', _tr('nom_complet', d.get('NOM'), d.get('PRENOM')), 'nom du candidat')
    jury = 0
    for i in range(1, NB_JURY + 1):
        if txt(d.get(f'J{i}_NOM')) and not txt(d.get(f'J{i}_NOM_AR')):
            a = _tr('nom_complet', d.get(f'J{i}_NOM'), d.get(f'J{i}_PRENOM'))
            if a: d[f'J{i}_NOM_AR'] = a; jury += 1
    if jury: tr.append(f'{jury} nom(s) de jury')
    if 'NOMS_JURY_AR' in d or any(txt(d.get(f'J{i}_NOM')) for i in range(1, NB_JURY + 1)):
        noms = [i for i in range(1, NB_JURY + 1) if txt(d.get(f'J{i}_NOM'))]
        d['NOMS_JURY_AR'] = 'ok' if noms and all(txt(d.get(f'J{i}_NOM_AR')) for i in noms) else ''
    d['TRADUIT_AR'] = ', '.join(tr)
    return d


# ------------------------------------------------------------------ export du formulaire -> format interne
ALIAS = {
    'NOM_EN_ARABE': 'NOM_AR', 'PRENOM_EN_ARABE': 'PRENOM_AR',
    'LIEU_DE_NAISSANCE': 'LIEU_NAISSANCE', 'LIEU_DE_NAISSANCE_EN_ARABE': 'LIEU_NAISSANCE_AR', 'DATE_DE_NAISSANCE': 'DATE_NAISSANCE',
    'DATE_DE_SOUTENANCE': 'DATE_SOUTENANCE', 'LIEU_DE_SOUTENANCE': 'LIEU_SOUTENANCE', 'HEURE_DE_SOUTENANCE': 'HEURE',
    'NOM_COMPLET_DE_L_ENCADRANT': 'ENCADRANT', 'NOM_COMPLET_DE_L_ENCADRANT_EN_ARABE': 'ENCADRANT_AR',
    'TITRE_DE_LA_THESE': 'TITRE', 'TITRE_DE_LA_THESE_EN_ARABE': 'TITRE_AR',
    'NOUVEAU_TITRE_DE_LA_THESE_SI_MODIFIE_PAR_LE_JURY': 'NOUVEAU_TITRE', 'TYPE_DE_DOSSIER': 'TYPE',
    'FORMATION_DOCTORAL': 'FORMATION', 'FORMATION_DOCTORALE': 'FORMATION',
    'FORMATION_DOCTORAL_EN_ARABE': 'FORMATION_AR', 'FORMATION_DOCTORALE_EN_ARABE': 'FORMATION_AR',
    'DISCIPLINE_EN_ARABE': 'DISCIPLINE_AR', 'SPECIALITE_EN_ARABE': 'SPECIALITE_AR', 'MENTION_EN_ARABE': 'MENTION_AR',
    'SPETIALITE': 'SPECIALITE', 'SPETIALITE_EN_ARABE': 'SPECIALITE_AR',          # colonne « Spétialité » du formulaire
    'CED_EN_ARABE': 'CED_AR', 'ADRESSE_E_MAIL': 'EMAIL', 'NUMERO_DE_TELEPHONE': 'TELEPHONE'}
JURY_COL = re.compile(r'^(NOM|PRENOM|GRADE|ROLE)_DU_MEMBRE_DU_JURY_N_(\d+)(_EN_ARABE)?$')
JURY_ETAB = re.compile(r'^ETABLISSEMENT_D_AFFILIATION_DU_MEMBRE_DU_JURY_N_(\d+)')
ARABE = re.compile(r'[\u0600-\u06FF]')
PRENOMS_M = set('''mohamed mohammed ahmed adnane abdelghani abdellah abdelaziz abderrahim abderrahmane abdelhamid abdelkader abdelilah
 abdessamad abdelouahed aziz azidine said hassan hicham mustapha mounir karim omar youssef yassine zohayr zouhair anas hamza adil redouan
 ali lahcen brahim ibrahim khalid khalil mehdi nabil noureddine samir soufiane tarik taoufik younes yahya ayoub amine othmane driss ismail
 jamal kamal larbi lahoucine hafid mourad hakim nouh david yannick rachid abdelhak abdelmajid aziz bilal fouad hamid hassane jaouad
 lahoussine mohcine moulay najib omar reda rachid redouane said slimane tarek yassir zakaria
 mossab mosaab musab mohssine mohsine abdelali abdeljalil abdelkrim abdelatif abdellatif abdennour abdeslam
 hassan hosni hussein houssam ilyas ilias ismael jalal marouane marwane mouad mouhcine nizar oussama ossama
 rida saad salah sami soufian taha walid wassim yassir youness zakariae zakaria aymane ayman badr hatim'''.split())
PRENOMS_F = set('''imane fatima zineb soumia fouzia salma khadija samira latifa meryem mariam maryam nadia asmae asmaa hafsa houda hanane hajar
 ikram ilham kawtar kaoutar loubna nawal naima nezha rkia rachida saida sara sanaa siham souad wafae wafa zahra amal aicha btissam btissame
 chaimae chaimaa bouchra hind jamila karima laila leila malika mouna mounia oumaima rajae sabah safae soukaina yasmine altaf lamia hasna
 fadwa ghizlane hakima hayat nabila najat sabrina sofia widad zakia rizlane rizlan soumaya soumiya mahassine mahassin
 fatimazahra fatimaezzahra nouhaila hiba khaoula ouafae ghita hind ibtissam ikrame kenza manal maha meriem nora oumayma rim
 sanae selma salima sihame yousra zahira zainab zineb chaima chaymae asma
 meriam meryam mariem myriam fatimazohra fatiha naoual nisrine nissrine noura nour ouiam ouissal rania rajaa
 sakina samia sanaa soukaina touria wiam yasmina zohra halima habiba hasnae ihssane ihsane imene jihane'''.split())


def collapse(s): return re.sub(r'\s+', ' ', txt(s))
def sans_accents(s): return ''.join(c for c in unicodedata.normalize('NFD', txt(s)) if unicodedata.category(c) != 'Mn')
def nom_maj(s): return collapse(s).upper()
def prenom_propre(s):
    cap = lambda w: w[:1].upper() + w[1:].lower()
    return ' '.join('-'.join(cap(x) for x in mot.split('-')) for mot in collapse(s).split(' ') if mot)
def casse_propre(s):                       # « AGADIR » / « agadir » -> « Agadir » ; texte mixte inchangé
    s = collapse(s)
    return prenom_propre(s) if s and (s.isupper() or s.islower()) else s
def mention_propre(s):                     # « trés honorable » / « TRES HONORABLE » -> « Très Honorable »
    m = sans_accents(collapse(s)).lower()
    return {'tres honorable': 'Très Honorable', 'honorable': 'Honorable'}.get(m, collapse(s))
def deviner_sexe(prenom):
    p = sans_accents(collapse(prenom)).lower().replace('-', ' ').split()
    p = p[0] if p else ''
    return 'M' if p in PRENOMS_M else 'F' if p in PRENOMS_F else ''
def meme_personne(enc, nom, prenom):
    E = set(sans_accents(enc).lower().split()); N = set(sans_accents(nom).lower().split()); P = sans_accents(prenom).lower().split()
    return bool(N) and N <= E and (not P or P[0] in E)
ORDRE_ROLES = ['president', 'directeur', 'co-directeur', 'rapporteur', 'examinateur', 'membre']
def role_fr(raw):
    r = sans_accents(raw).lower()
    if not r: return ''
    if 'president' in r:
        pres = 'Présidente' if 'presidente' in r else 'Président'
        return f'Rapporteur et {pres}' if 'rapporteur' in r else pres
    if 'co-direct' in r or 'codirect' in r: return 'Co-directrice' if 'trice' in r else 'Co-directeur'
    if 'direct' in r: return 'Directrice de thèse' if 'trice' in r else 'Directeur de thèse'
    if 'rapporteur' in r: return 'Rapporteur'
    if 'examinat' in r: return 'Examinateur'
    return 'Membre' if 'membre' in r else collapse(raw)
def rang_role(role):
    r = sans_accents(role).lower()
    if 'president' in r: return 0
    for i, k in enumerate(ORDRE_ROLES):
        if r.startswith(k) or (k == 'directeur' and r.startswith('directrice')) or (k == 'co-directeur' and r.startswith('co-directrice')): return i
    return len(ORDRE_ROLES)
def role_ar(role, i=0, sexe=''):
    """Qualité du membre en arabe : toujours au masculin (رئيسا / مشرفا / عضوا), quel que soit le genre.
    Seul le grade (grade_ar) est accordé au féminin."""
    r = sans_accents(role).lower()
    if not r: r = 'president' if i == 1 else 'directeur' if i == 2 else 'membre'
    if 'president' in r: return 'رئيسا'
    if r.startswith('co-direct'): return 'مشرفا مشاركا'
    if r.startswith('direct'): return 'مشرفا'
    return 'عضوا'


def depuis_formulaire(d):
    """Une ligne de l'export du formulaire -> dict au format interne (NOM, PRENOM, J1_NOM, J1_AFFILIATION, ...)."""
    out, jury = {}, {}
    for k, v in d.items():
        if v in (None, ''): continue
        m = JURY_COL.match(k)
        if m:
            jury.setdefault(int(m.group(2)), {})[m.group(1) + ('_AR' if m.group(3) else '')] = v; continue
        m = JURY_ETAB.match(k)
        if m: jury.setdefault(int(m.group(1)), {})['ETAB'] = v; continue
        out[ALIAS.get(k, k)] = v
    out['NOM'], out['PRENOM'] = nom_maj(out.get('NOM')), prenom_propre(out.get('PRENOM'))
    cne = out.get('CNE')
    out['CNE'] = txt(int(cne) if isinstance(cne, float) and cne == int(cne) else cne).replace(' ', '')
    out['CIN'] = txt(out.get('CIN')).replace(' ', '').upper()
    for k in ('NOM_AR', 'PRENOM_AR', 'LIEU_NAISSANCE_AR', 'SPECIALITE', 'LIEU_SOUTENANCE', 'ENCADRANT'): out[k] = collapse(out.get(k))
    out['LIEU_NAISSANCE'] = casse_propre(out.get('LIEU_NAISSANCE'))
    out['DISCIPLINE'] = casse_propre(out.get('DISCIPLINE'))
    out['MENTION'] = mention_propre(out.get('MENTION'))
    if not txt(out.get('DISCIPLINE_AR')): out['DISCIPLINE_AR'] = DISCIPLINE_AR.get(sans_accents(out['DISCIPLINE']).lower(), '')
    out['TITRE'] = collapse(out.get('NOUVEAU_TITRE') or out.get('TITRE')).rstrip('.')
    if sans_accents(out.get('CED')).lower().startswith('centre'): out['CED'] = ''     # « Centre d'Études Doctorales » = pas le nom du CED
    if not txt(out.get('SEXE')):
        out['SEXE'] = deviner_sexe(out['PRENOM']); out['SEXE_DEDUIT'] = 'oui' if out['SEXE'] else ''
    # --- jury
    membres = []
    for i in sorted(jury):
        j = jury[i]; nom, pre, pre_ar = nom_maj(j.get('NOM')), prenom_propre(j.get('PRENOM')), collapse(j.get('PRENOM_AR'))
        if not nom: continue
        if not pre and pre_ar and not ARABE.search(pre_ar): pre, pre_ar = prenom_propre(pre_ar), ''   # prénom latin saisi dans la case arabe
        elif ARABE.search(pre) and not ARABE.search(pre_ar):                                        # prénom arabe saisi dans la case latine
            pre, pre_ar = prenom_propre(pre_ar), collapse(j.get('PRENOM'))
        raw = txt(j.get('GRADE')); code = (raw.split(':')[0] if ':' in raw else raw).strip().upper()
        membres.append(dict(nom=nom, pre=pre, ar=' '.join(x for x in (collapse(j.get('NOM_AR')), pre_ar) if x),
                            code=code if code in GRADE_FR else '', etab=collapse(j.get('ETAB')), role=role_fr(txt(j.get('ROLE')))))
    roles = [m['role'] for m in membres]
    if membres and all(roles):
        membres.sort(key=lambda m: rang_role(m['role']))                      # président, directeur, rapporteurs, examinateurs...
    elif membres and not any(roles):
        for m in membres:
            if meme_personne(out['ENCADRANT'], m['nom'], m['pre']): m['role'] = 'Directeur de thèse'; break
        if not membres[0]['role']: membres[0]['role'] = 'Président'      # convention des PV : le 1er membre préside
    for n, m in enumerate(membres[:NB_JURY], 1):
        out[f'J{n}_NOM'], out[f'J{n}_PRENOM'], out[f'J{n}_NOM_AR'], out[f'J{n}_ROLE'] = m['nom'], m['pre'], m['ar'], m['role']
        out[f'J{n}_AFFILIATION'] = ', '.join(x for x in (m['code'], m['etab']) if x)
        out[f'J{n}_SEXE'] = deviner_sexe(m['pre'])
    out['ROLES_JURY'] = 'ok' if membres and all(m['role'] for m in membres) else ''
    out['NOMS_JURY_AR'] = 'ok' if membres and all(m['ar'] for m in membres) else ''
    return out


def valeurs(r):
    d = r.get('DATE_SOUTENANCE'); h = r.get('HEURE')
    if isinstance(h, (dt.time, dt.datetime)): h = f"{h.hour}h{h.minute:02d}"
    long_ = txt(d)
    if isinstance(d, (dt.date, dt.datetime)):
        long_ = f"{JOURS[d.weekday()]} {d.day:02d} {MOIS[d.month-1]} {d.year}" + (f" à {txt(h)}" if h else '')
    sx = txt(r.get('SEXE')).upper()[:1]; homme, femme = sx == 'M', sx == 'F'
    v = {k: txt(x) for k, x in r.items()}
    v.update(DATE_NAISSANCE=date_fr(r.get('DATE_NAISSANCE')), DATE_SOUTENANCE=date_fr(d), DATE_LONGUE=long_,
             CIVILITE='M.' if homme else 'Mme' if femme else 'M./Mme',
             CIVILITE_LONGUE='Monsieur' if homme else 'Madame' if femme else 'Madame/Monsieur',
             NE='né' if homme else 'née' if femme else 'né(e)',
             ANNEE=str(d.year) if isinstance(d, (dt.date, dt.datetime)) else '')
    # --- accord masculin / féminin du diplôme (sexe inconnu -> masculin, signalé dans RAPPORT.txt)
    v.update(ETUDIANT='l’étudiante' if femme else 'l’étudiant',
             NE_DIPLOME='née' if femme else 'né',
             ETUDIANT_AR='الطالبة' if femme else 'الطالب',
             NE_AR='المزدادة' if femme else 'المزداد',
             OBTENU_AR='أحرزت' if femme else 'أحرز')
    for i in range(1, NB_JURY + 1):
        n, p = txt(r.get(f'J{i}_NOM')), txt(r.get(f'J{i}_PRENOM'))
        v[f'J{i}_COURT'] = f"{p[:1]}. {n}" if n else ''
        v[f'J{i}_AFF'] = txt(r.get(f'J{i}_AFFILIATION'))
    v.update(NOM_COMPLET=f"{v.get('NOM','')} {v.get('PRENOM','')}".strip(),
             ETABLISSEMENT=v.get('ETABLISSEMENT') or ETAB_DEFAUT,
             ETABLISSEMENT_AR=v.get('ETABLISSEMENT_AR') or ETAB_AR_DEFAUT,
             LIEU_REDACTION_AR=v.get('LIEU_REDACTION_AR') or 'أكادير',
             DATE_SOUT_AR=date_ar(d), DATE_NAISSANCE_AR=date_ar(r.get('DATE_NAISSANCE')),
             MENTION_AR=v.get('MENTION_AR') or MENTION_AR.get(txt(r.get('MENTION')).lower(), ''),
             FORMATION_AR=v.get('FORMATION_AR') or FORMATION_AR.get(txt(r.get('FORMATION')).lower(), ''))
    # --- genre des membres du jury (Excel > rôle > listes locales > API genderize.io), pour l'arabe du diplôme
    avec_nom = [i for i in range(1, NB_JURY + 1) if txt(r.get(f'J{i}_NOM'))]
    res = genre.sexes_jury([(r.get(f'J{i}_PRENOM'), r.get(f'J{i}_SEXE'), r.get(f'J{i}_ROLE')) for i in avec_nom], deviner_sexe)
    sexe_j = {i: s for i, (s, _) in zip(avec_nom, res)}
    src = [s for _, s in res]
    v['_genre_jury'] = (f"{src.count('api')} via API" if 'api' in src else '',
                        ', '.join(f"{txt(r.get(f'J{i}_PRENOM'))} {txt(r.get(f'J{i}_NOM'))}" for i, (_, s) in zip(avec_nom, res) if not s))
    lignes = []
    for i in range(1, 8):
        na = txt(r.get(f'J{i}_NOM_AR'))
        if not na: continue
        grade = grade_ar(txt(r.get(f'J{i}_AFFILIATION')).split(',')[0].strip(), sexe_j.get(i, ''))
        role = role_ar('', i, sexe_j.get(i, ''))
        lignes.append(f"{na}، {grade} / {role}")
    for k in range(3): v[f'JURY_AR_{k+1}'] = lignes[k] if k < len(lignes) else BLANC
    v['_extras'] = lignes[3:]
    for k in ('NUM_PV', 'SPECIALITE', 'MENTION', 'NOM_AR', 'LIEU_NAISSANCE_AR', 'DISCIPLINE_AR', 'SPECIALITE_AR', 'FORMATION_AR', 'TITRE_AR', 'MENTION_AR', 'LIEU_NAISSANCE', 'LIEU_SOUTENANCE', 'DISCIPLINE', 'CIN', 'CNE', 'FORMATION', 'DISCIPLINE_AR', 'DATE_NAISSANCE', 'DATE_SOUTENANCE', 'TITRE'):
        v[k] = v.get(k) or BLANC

    # --- Champs des modèles Attestation.docx et Diplome.docx (faits d'après les exemples OUBRAI) ---
    auj = dt.date.today()
    v.update(DATE_EDITION=auj.strftime('%d/%m/%Y'), DATE_EDITION_LONG=date_longue(auj), DATE_EDITION_AR=date_ar(auj),
             DATE_SOUTENANCE_LONG=date_longue(d) or BLANC, DATE_NAISSANCE_LONG=date_longue(r.get('DATE_NAISSANCE')) or BLANC)
    v['DATE_SOUT_AR'] = date_ar(d) or BLANC
    v['DATE_NAISSANCE_AR'] = date_ar(r.get('DATE_NAISSANCE')) or BLANC
    v['NOM_AR_COMPLET'] = ' '.join(x for x in (txt(r.get('NOM_AR')), txt(r.get('PRENOM_AR'))) if x) or BLANC
    v['CED'] = txt(r.get('CED')) or CED_DEFAUT
    if sans_accents(v['CED']).lower() == sans_accents(CED_DEFAUT).lower(): v['CED'] = CED_DEFAUT      # « Sciences Et Techniques… » -> graphie officielle
    v['CED_AR'] = txt(r.get('CED_AR')) or (CED_AR_DEFAUT if sans_accents(v['CED']).lower() == sans_accents(CED_DEFAUT).lower() else BLANC)
    v['DISCIPLINE_AR'] = txt(r.get('DISCIPLINE_AR')) or DISCIPLINE_AR.get(sans_accents(r.get('DISCIPLINE')).lower(), '') or BLANC
    v['BAR_CNE'], v['BAR_CIN'] = repere('BAR_CNE'), repere('BAR_CIN')   # remplacés par des images après fill()
    v['_codes'] = {'BAR_CNE': r.get('CNE'), 'BAR_CIN': r.get('CIN')}
    roles_saisis = any(txt(r.get(f'J{i}_ROLE')) for i in range(1, NB_JURY + 1))
    for i in range(1, NB_JURY + 1):
        n, p = txt(r.get(f'J{i}_NOM')), txt(r.get(f'J{i}_PRENOM'))
        parts = [x.strip() for x in txt(r.get(f'J{i}_AFFILIATION')).split(',') if x.strip()]
        code = parts[0] if parts and parts[0] in GRADE_FR else ''
        reste = parts[1:] if code else parts
        role = txt(r.get(f'J{i}_ROLE')) or ('' if roles_saisis else 'Président' if i == 1 else 'Directeur de thèse' if i == 2 else '')
        v[f'JURY_{i}'] = 'x' if n else ''                                   # marqueur {{?JURY_i}} : ligne supprimée si vide
        v[f'JURY_{i}_NOM_GRADE'] = ((f"{p} {n}".strip() + (f", {GRADE_FR[code]}" if code else '')) + ',') if n else ''
        v[f'JURY_{i}_ETABLISSEMENT'] = ', '.join(reste) if n else ''
        v[f'JURY_{i}_ROLE'] = role
        v[f'JURY_{i}_NOM_AR'] = txt(r.get(f'J{i}_NOM_AR')) or BLANC if n else ''
        sj = sexe_j.get(i, '')
        v[f'JURY_{i}_SEXE'] = sj
        v[f'JURY_{i}_GRADE_AR'] = f"{grade_ar(code, sj)} ،" if code in GRADE_AR else ''
        v[f'JURY_{i}_GRADE_ETAB_AR'] = (' ، '.join(x for x in (grade_ar(code, sj), txt(r.get(f'J{i}_ETAB_AR')) or etab_ar(reste)) if x)) if n else ''
        v[f'JURY_{i}_ROLE_AR'] = role_ar(role, i, sj) if n else ''
    return v


REQ_PV = ['NOM','PRENOM','DATE_NAISSANCE','LIEU_NAISSANCE','CNE','CIN','FORMATION','DISCIPLINE','SPECIALITE','TITRE','DATE_SOUTENANCE','HEURE','LIEU_SOUTENANCE','NUM_PV','J1_NOM','J2_NOM']
REQ_ATT = ['NOM','PRENOM','SEXE','DATE_NAISSANCE','LIEU_NAISSANCE','CNE','CIN','FORMATION','DISCIPLINE','SPECIALITE','TITRE','DATE_SOUTENANCE','MENTION','J1_NOM','ROLES_JURY']
REQ_ATT_HU = ['NOM','PRENOM','SEXE','DATE_NAISSANCE','LIEU_NAISSANCE','CNE','CIN','DISCIPLINE','SPECIALITE','DATE_SOUTENANCE','J1_NOM','ROLES_JURY']
REQ_DIP_HU = ['NOM','PRENOM','NOM_AR','DATE_NAISSANCE','CIN','SPECIALITE','SPECIALITE_AR','DATE_SOUTENANCE','J1_NOM','ROLES_JURY','NOMS_JURY_AR']
REQ_INV = ['NOM','PRENOM','SEXE','TITRE','DATE_SOUTENANCE','HEURE','LIEU_SOUTENANCE','J1_NOM','J1_PRENOM','J1_AFFILIATION','J2_NOM','J2_PRENOM','J2_AFFILIATION']
REQ_RES = ['NOM','PRENOM','CIN','CNE','TITRE','DATE_SOUTENANCE','FORMATION','J1_NOM','J1_ROLE']
REQ_DIP = ['NOM','PRENOM','NOM_AR','DATE_NAISSANCE','LIEU_NAISSANCE','LIEU_NAISSANCE_AR','CNE','CIN','DISCIPLINE','DISCIPLINE_AR','SPECIALITE','SPECIALITE_AR','TITRE','TITRE_AR','MENTION','DATE_SOUTENANCE','J1_NOM','ROLES_JURY','NOMS_JURY_AR']


def dispo(prog, doc):
    f = {'pv': ['Page_0.docx','Page_1.docx','Page_2.docx'], 'diplome': ['Diplome.docx'], 'attestation': ['Attestation.docx'],
         'invitation': ['Invitation.docx'], 'resume': ['Resume.docx']}[doc]
    return all(os.path.exists(os.path.join(MOD, prog, x)) for x in f)


GRADE_FR = {'PES': "Professeur de l'Enseignement Supérieur", 'PH': 'Professeur Habilité', 'MCH': 'Maître de Conférences Habilité', 'MC': 'Maître de Conférences', 'P': 'Professeur'}


# ------------------------------------------------------------------ Résumé des avis des rapporteurs (Resume.docx)
NB_DIR_RES, NB_RAP_RES = 3, 6
SIGLES_ETAB = {     # établissement (minuscules, sans accents) -> sigle, comme sur le modèle (« FS, Agadir »)
    'faculte des sciences': 'FS', 'faculte des sciences et techniques': 'FST', 'faculte polydisciplinaire': 'FP',
    'ecole superieure de technologie': 'EST', 'ecole nationale des sciences appliquees': 'ENSA',
    'ecole nationale de commerce et de gestion': 'ENCG', 'faculte de medecine et de pharmacie': 'FMP',
    "ecole nationale superieure d'informatique et d'analyse des systemes": 'ENSIAS', 'ecole normale superieure': 'ENS',
    'faculte des lettres et des sciences humaines': 'FLSH', 'ecole nationale superieure des arts et metiers': 'ENSAM',
    'faculte des sciences juridiques, economiques et sociales': 'FSJES', 'faculte des sciences juridiques economiques et sociales': 'FSJES',
    'institut agronomique et veterinaire hassan ii': 'IAV', 'faculte des sciences appliquees': 'FSA',
    'faculte des sciences de l\'education': 'FSE', 'ecole nationale d\'agriculture': 'ENA'}
MOTS_VIDES = {'de', 'la', 'le', 'les', 'des', 'du', 'et', 'd', 'l', 'en', 'aux', 'au', 'a'}


def _cle_etab(s):
    """Clé tolérante aux fautes du formulaire : « Fcaulté de Sciences » -> 'faculte des sciences'."""
    s = re.sub(r'\s+', ' ', sans_accents(s).lower().replace('’', "'")).strip(' .')
    s = re.sub(r'\b(fcaulte|faculte|facultee|fac)\b', 'faculte', s)
    s = re.sub(r'\b(faculte|ecole[\w ]*?) de sciences\b', r'\1 des sciences', s)
    return re.sub(r'\s*\(.*?\)', '', s)                              # « École … (EST) » -> sans le sigle entre parenthèses


def etab_court(reste, virgule=True):
    """['Faculté des Sciences', 'Agadir'] -> 'FS, Agadir' ; ['Faculté des Sciences', 'UIZ', 'Agadir'] -> 'FS, Agadir' (ville = dernière partie).
    Texte d'origine si l'établissement n'a pas de sigle connu."""
    sep = ', ' if virgule else ' '
    if len(reste) >= 2:
        for n in range(len(reste) - 1, 0, -1):                      # le plus long d'abord
            s = SIGLES_ETAB.get(_cle_etab(', '.join(reste[:n])))
            if s: return s + sep + reste[-1]
    elif reste and SIGLES_ETAB.get(_cle_etab(reste[0])):
        return SIGLES_ETAB[_cle_etab(reste[0])]
    return ', '.join(reste)


def sigle_formation(r):
    """Colonne « Sigle formation » si elle existe, sinon initiales : « Sciences de la Vie et Ressources Naturelles » -> SVRN."""
    for k in ('FORMATION_SIGLE', 'SIGLE_FORMATION', 'SIGLE_DE_LA_FORMATION', 'SIGLE'):
        if txt(r.get(k)): return txt(r.get(k)), False
    f = txt(r.get('FORMATION'))
    mots = [m for m in re.split(r"[\s,&'’/\-]+", sans_accents(f)) if m and m.lower() not in MOTS_VIDES]
    return (''.join(m[0].upper() for m in mots), True) if len(mots) >= 3 else (f, False)


def annee_inscription(r):
    """« 2021/2022 » ; une date (01/10/2021) ou une année (2021) donnent l'année universitaire."""
    for k in ('ANNEE_INSCRIPTION', 'ANNEE_D_INSCRIPTION', 'ANNEE_DE_PREMIERE_INSCRIPTION', 'ANNEE_DE_1ERE_INSCRIPTION',
              'ANNEE_UNIVERSITAIRE_D_INSCRIPTION', 'DATE_D_INSCRIPTION', 'DATE_INSCRIPTION', 'DATE_DE_PREMIERE_INSCRIPTION'):
        x = r.get(k)
        if isinstance(x, (dt.date, dt.datetime)):
            a = x.year if x.month >= 9 else x.year - 1
            return f'{a}/{a + 1}'
        if isinstance(x, float) and x == int(x): x = int(x)
        if isinstance(x, int) and 1990 < x < 2100: return f'{x}/{x + 1}'
        if txt(x): return txt(x)
    return ''


def resume_valeurs(r, v):
    """Champs propres à Resume.docx (candidat : RES_* ; directeurs : DIR_k_* ; rapporteurs : RAP_k_*) + notes pour RAPPORT.txt."""
    notes = []
    sigle, deduit = sigle_formation(r)
    if deduit: notes.append(f'formation doctorale abrégée en « {sigle} » : à vérifier')
    w = dict(RES_NOM=', '.join(x for x in (txt(r.get('NOM')), txt(r.get('PRENOM'))) if x) or BLANC,
             RES_TYPE=txt(r.get('TYPE_DOCTORAT')) or 'DCED',
             RES_ANNEE_INSCRIPTION=annee_inscription(r) or BLANC,
             RES_FORMATION=sigle or BLANC)
    if w['RES_ANNEE_INSCRIPTION'] == BLANC: notes.append("année d'inscription absente")
    disc_cand = txt(r.get('DISCIPLINE'))                               # discipline d'un membre inconnue : celle du candidat
    dirs, raps, avis_defaut = [], [], 0
    for i in range(1, NB_JURY + 1):
        n, p = txt(r.get(f'J{i}_NOM')), txt(r.get(f'J{i}_PRENOM'))
        if not n: continue
        role = sans_accents(v.get(f'JURY_{i}_ROLE')).lower()
        parts = [x.strip() for x in txt(r.get(f'J{i}_AFFILIATION')).split(',') if x.strip()]
        code = parts[0] if parts and parts[0] in GRADE_FR else ''
        reste = parts[1:] if code else parts
        disc = txt(r.get(f'J{i}_DISCIPLINE')) or txt(r.get(f'J{i}_SPECIALITE')) or disc_cand
        if role.startswith(('direct', 'co-direct', 'codirect')):
            dirs.append(dict(NOM=f'{n}, {p[:1]}.' if p else n, GRADE=code, DISC=disc, ETAB=etab_court(reste)))
        if 'rapporteur' in role:
            avis = txt(r.get(f'J{i}_AVIS'))
            if not avis: avis = 'Favorable'; avis_defaut += 1
            raps.append(dict(NOM=f'{n}, {p[:1]}.' if p else n, GRADE=code, DISC=disc, ETAB=etab_court(reste), AVIS=avis))
    if not dirs and txt(r.get('ENCADRANT')):                          # directeur absent du jury : on reprend l'encadrant
        dirs.append(dict(NOM=txt(r.get('ENCADRANT')), GRADE='', DISC=disc_cand, ETAB=''))
    for nom, liste, nb in (('DIR', dirs, NB_DIR_RES), ('RAP', raps, NB_RAP_RES)):
        for k in range(1, nb + 1):
            m = liste[k - 1] if k <= len(liste) else None
            w[f'{nom}_{k}'] = 'x' if m else ''                            # {{?DIR_k}} / {{?RAP_k}} : ligne supprimée si vide
            for c in ('NOM', 'GRADE', 'DISC', 'ETAB', 'AVIS'): w[f'{nom}_{k}_{c}'] = (m or {}).get(c, '')
        if len(liste) > nb: notes.append(f'{len(liste) - nb} {nom.lower()} non affiché(s) (max {nb})')
    if not dirs: notes.append('aucun directeur de thèse')
    if not raps: notes.append('aucun rapporteur (rôles du jury absents ?)')
    if avis_defaut: notes.append(f'avis « Favorable » par défaut pour {avis_defaut} rapporteur(s)')
    return dict(v, **w), notes


def invitations(r, v, prog):
    """Une série de valeurs par membre du jury : J1 = président, tous les autres 'participer'.
    Colonnes optionnelles J{i}_SEXE (M/F, par défaut M)."""
    hu = prog == 'habilitation'
    out = []
    for i in range(1, 8):
        if not v.get(f'J{i}_NOM') or v[f'J{i}_NOM'] == BLANC: continue
        parts = [x.strip() for x in txt(r.get(f'J{i}_AFFILIATION')).split(',') if x.strip()]
        grade = GRADE_FR.get(parts[0], '') if parts else ''
        reste = parts[1:] if grade else parts
        inst = ', '.join(reste[:-1]) if len(reste) > 1 else (reste[0] if reste else '')
        ville = reste[-1].upper() if len(reste) > 1 else ''
        f = txt(r.get(f'J{i}_SEXE')).upper().startswith('F')
        appel = 'Madame' if f else 'Monsieur'; pres = i == 1
        court = v[f'J{i}_COURT']
        dest = f"{appel} le professeur {court}" if pres else f"{appel} {court}"
        lignes = [dest] + ([] if pres or not grade else [grade]) + ([inst] if inst else [])
        w = dict(v, INV_DEST='\n'.join(lignes), INV_VILLE=ville, INV_APPEL=appel,
                 INV_OBJET=f"Invitation pour {'présider' if pres else 'participer à'} un jury {'d’habilitation universitaire' if hu else 'de thèse'}",
                 INV_VERBE='présider le' if pres else 'participer au',
                 INV_NATURE='Jury d’Habilitation universitaire de' if hu else 'Jury de thèse de Doctorat de')
        out.append((i, propre(f"{v[f'J{i}_NOM']}_{txt(r.get(f'J{i}_PRENOM'))}"), w))
    return out


def avec_codes_barres(docx, v):
    """Remplace [[BAR_CNE]] / [[BAR_CIN]] du document rempli par les images Code 39 du CNE et de la CIN."""
    return inserer_codes_barres(docx, v['_codes'])[0]


def manquants(r, req): return [k for k in req if not txt(r.get(k))]


def prog_of(r):
    return 'habilitation' if txt(r.get('TYPE')).lower().startswith('h') else 'doctorat'


@app.route('/')
def index():
    return render_template('index.html', dispo={p: {d: dispo(p, d) for d in DOCS_VALIDES} for p in PROGRAMMES}, programmes=PROGRAMMES)


@app.post('/api/upload')
def upload():
    f = request.files.get('fichier')
    if not f or not f.filename.lower().endswith('.xlsx'): return jsonify(erreur='Fichier .xlsx requis'), 400
    token = uuid.uuid4().hex; path = os.path.join(UP, token + '.xlsx'); f.save(path)
    try:
        rows = lire(path)
    except Exception as e:
        os.remove(path)
        return jsonify(erreur=f'Lecture impossible : {e}'), 400
    if not rows:
        os.remove(path)
        return jsonify(erreur='Aucune ligne de données dans le fichier'), 400
    return jsonify(token=token, candidats=[dict(i=i, type=prog_of(r), nom=txt(r.get('NOM')), prenom=txt(r.get('PRENOM')), date=date_fr(r.get('DATE_SOUTENANCE')),
        manque_pv=manquants(r, REQ_PV), manque_diplome=manquants(r, REQ_DIP_HU if prog_of(r) == 'habilitation' else REQ_DIP), manque_attestation=manquants(r, REQ_ATT_HU if prog_of(r) == 'habilitation' else REQ_ATT), manque_invitation=manquants(r, REQ_INV), manque_resume=manquants(r, REQ_RES)) for i, r in enumerate(rows)])


@app.post('/api/generer')
def generer():
    d = request.get_json(force=True, silent=True) or {}
    docs = [x for x in (d.get('docs') or []) if x in DOCS_VALIDES]
    idx = []                                    # le JSON peut envoyer "3" au lieu de 3
    for i in (d.get('indices') or []):
        try: idx.append(int(i))
        except (TypeError, ValueError): pass
    idx = list(dict.fromkeys(idx))              # sans doublons, ordre conservé
    token = d.get('token')
    if not token or not idx or not docs: return jsonify(erreur='Paramètres invalides'), 400
    path = os.path.join(UP, os.path.basename(str(token)) + '.xlsx')
    if not os.path.exists(path): return jsonify(erreur='Fichier introuvable, ré-importez le fichier Excel'), 400

    rows = lire(path); buf = io.BytesIO(); rapport = []
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
        for i in idx:
            if not 0 <= i < len(rows):
                rapport.append(f"IGNORÉ : indice {i} hors limites"); continue
            r = rows[i]
            try:
                v = valeurs(r); prog = prog_of(r)
                nom = propre(f"{v.get('NOM','')}_{v.get('PRENOM','')}"); base = f"{prog}/{i+1:02d}_{nom}/"
            except Exception as e:
                rapport.append(f"ERREUR : ligne {i+1} - {type(e).__name__} {e} | colonnes lues : {', '.join(r.keys())}"); continue
            qui = f"{v.get('NOM','')} {v.get('PRENOM','')} ({PROGRAMMES[prog]})"
            for doc in docs:
                if not dispo(prog, doc):
                    rapport.append(f"IGNORÉ : {qui} - modèle de {doc} indisponible"); continue
                try:
                    if doc == 'pv':
                        for n in range(3): z.writestr(f'{base}Page_{n}_PV_{nom}.docx', fill(os.path.join(MOD, prog, f'Page_{n}.docx'), v))
                    elif doc == 'invitation':
                        for k, jn, w in invitations(r, v, prog): z.writestr(f'{base}Invitations/Invitation_J{k}_{jn}.docx', fill(os.path.join(MOD, prog, 'Invitation.docx'), w))
                    elif doc == 'attestation':
                        z.writestr(f'{base}Attestation_{nom}.docx', avec_codes_barres(fill(os.path.join(MOD, prog, 'Attestation.docx'), v), v))
                    elif doc == 'resume':
                        w, notes_res = resume_valeurs(r, v)
                        z.writestr(f'{base}Resume_{nom}.docx', fill(os.path.join(MOD, prog, 'Resume.docx'), w))
                    else:
                        z.writestr(f'{base}Diplome_{nom}.docx', aligner(accorder(avec_codes_barres(fill(os.path.join(MOD, prog, 'Diplome.docx'), v), v), r.get('SEXE'))))
                    sexe = txt(r.get('SEXE')).upper()[:1]
                    note_sexe = ('' if doc != 'diplome' else ' [sexe inconnu : diplôme au masculin, à vérifier]' if sexe not in ('M', 'F')
                                 else f" [{'femme' if sexe == 'F' else 'homme'}, déduit du prénom : à vérifier]" if r.get('SEXE_DEDUIT') else '')
                    api_n, inconnus = v.get('_genre_jury', ('', ''))
                    if doc == 'diplome':
                        note_sexe += (f" [genre du jury : {api_n}, à vérifier]" if api_n else '') + \
                                     (f" [genre du jury inconnu (masculin par défaut) : {inconnus}]" if inconnus else '')
                    if doc == 'resume' and notes_res: note_sexe += ' [' + ' ; '.join(notes_res) + ']'
                    note_spe = ' [spécialité absente : reprise de la discipline]' if r.get('SPECIALITE_DEDUITE') and doc in ('attestation', 'diplome', 'pv') else ''
                    rapport.append(f"OK : {qui} - {doc}" + (f" [arabe traduit automatiquement : {v['TRADUIT_AR']}]" if v.get('TRADUIT_AR') and doc in ('attestation', 'diplome') else '') + note_sexe + note_spe)
                except Exception as e:                      # une erreur ne doit pas casser tout le zip
                    rapport.append(f"ERREUR : {qui} - {doc} : {e}")
        z.writestr('RAPPORT.txt', '\n'.join(rapport))
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"{'_'.join(docs)}_{len(idx)}_dossiers.zip", mimetype='application/zip')


if __name__ == '__main__': app.run(host='0.0.0.0', port=5000)